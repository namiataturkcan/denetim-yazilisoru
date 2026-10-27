"""Deterministic incoming-paper reports from uploaded tables; no model or API calls."""
import csv
import io
import hashlib
import re
import zipfile
from contextlib import contextmanager
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
from xml.etree import ElementTree

from openpyxl import load_workbook

from recipient_validation import normalize


MAX_TABLE_BYTES = 12 * 1024 * 1024
MAX_ROWS = 100_000
MAX_COLUMNS = 128
MAX_UNPACKED_BYTES = 160 * 1024 * 1024
REPORT_VERSION = hashlib.sha256(Path(__file__).read_bytes()).hexdigest()[:12]
STATUS_LABELS = {
    'announced': 'Cevaplanmadığı Gelen Kağıtlarda İlan Edildi',
    'pending': '15 Günlük Cevaplanma Süresi Devam Ediyor',
    'answered': 'Cevaplandı',
    'withdrawn': 'Geri alındı',
    'unknown': 'Durumu tanınmayan',
    'conflict': 'Çelişkili kayıt',
}
WORKFLOW_STAGES = {
    'prepare': {'label': 'Memur işlemi · İlan edilecekler', 'group': 'pending',
                'count_label': 'İlan edilecek', 'ranges_label': 'İlan için esas numarası listesi'},
    'verify': {'label': 'Uzman kontrolü · İlan edilenler', 'group': 'announced',
               'count_label': 'İlan edilmiş', 'ranges_label': 'İlan edilmiş esas numarası listesi'},
}


class TableError(ValueError):
    pass


def key(value):
    text = normalize(str(value or '')).translate(str.maketrans('çğıöşüâû', 'cgiosuau'))
    return re.sub(r'[^a-z0-9]', '', text)


_status_groups = {
    key('Cevaplanmadığı Gelen Kağıtlarda İlan Edildi'): 'announced',
    key('15 Günlük Cevaplanma Süresi Devam Ediyor'): 'pending',
    key('Cevaplandı'): 'answered',
    key('Süresi İçinde Cevaplandı'): 'answered',
    key('Süresi Geçtikten Sonra Cevaplandı'): 'answered',
    key('Önerge Sahibi Tarafından Geri Alındı'): 'withdrawn',
    key('Geri Alındı'): 'withdrawn',
}
_column_aliases = {
    'number': ('Esas No', 'Esas Numarası', 'Esas Numara'),
    'status': ('Son Durum', 'Durum', 'Durumu', 'Statü', 'Önerge Durumu', 'Önergenin Son Durumu'),
    'recipient': ('Muhatap Unvanı', 'Muhatap Ünvanı', 'Muhatap', 'Unvan', 'Muhatap Unvan', 'Muhatap Ünvan'),
    'term': ('Donem Adı', 'Dönem', 'Dönem Adı', 'Yasama Dönemi'),
}
_header_keys = {field: {key(alias) for alias in aliases} for field, aliases in _column_aliases.items()}


def parse_number(value):
    if isinstance(value, bool) or value is None:
        raise ValueError('Geçersiz esas numarası')
    if isinstance(value, float):
        if not value.is_integer():
            raise ValueError('Geçersiz esas numarası')
        value = int(value)
    text = str(value).strip()
    text = re.sub(r'^7\s*/\s*', '', text)
    if re.fullmatch(r'[1-9]\d{0,2}(?:\.\d{3})+', text):
        text = text.replace('.', '')
    if not re.fullmatch(r'\d{1,9}', text) or int(text) <= 0:
        raise ValueError('Geçersiz esas numarası')
    return int(text)


def term_name(value):
    text = str(value or '').strip()
    match = re.fullmatch(r'(\d+)\s*\.?\s*(?:dönem)?', text, re.IGNORECASE)
    return f'{int(match[1])}. Dönem' if match else text


@dataclass(frozen=True, slots=True)
class Record:
    number: int | None
    status: str
    recipient: str
    term: str
    source_row: int


def read_records(rows):
    """Read named KKBS columns or the three leading columns of a headerless paste."""
    mapping = None
    records = []
    for row_number, row in enumerate(rows, 1):
        # Formatting can extend a sheet far beyond its actual data columns.
        row = list(row)
        while row and (row[-1] is None or not str(row[-1]).strip()):
            row.pop()
        if row_number > MAX_ROWS + 20:
            raise TableError('Tablo en fazla 100.000 veri satırı içerebilir.')
        if len(row) > MAX_COLUMNS:
            raise TableError('Tabloda en fazla 128 sütun destekleniyor.')
        if not any(value is not None and str(value).strip() for value in row):
            continue
        header = {}
        if mapping is None or (len(row) > mapping['number'] and key(row[mapping['number']]) in _header_keys['number']):
            normalized = [key(value) for value in row]
            for field, aliases in _column_aliases.items():
                matches = [i for i, column in enumerate(normalized) if column in _header_keys[field]]
                if len(matches) > 1:
                    raise TableError(f'Birden fazla {aliases[0]} sütunu bulundu; tek sütun kullanın.')
                if matches:
                    header[field] = matches[0]
        if {'number', 'status', 'recipient'} <= header.keys():
            mapping = header
            continue
        if {'number', 'status', 'recipient'} & header.keys():
            missing = [aliases[0] for field, aliases in _column_aliases.items()
                       if field in {'number', 'status', 'recipient'} and field not in header]
            raise TableError('Tablo başlıkları eksik veya tanınmıyor: ' + ', '.join(missing)
                             + '. Esas No, Son Durum ve Muhatap Unvanı başlıklarını birlikte ekleyin.')
        if mapping is None:
            # Historic ChatGPT examples were pasted without headers: number, status, title.
            try:
                parse_number(row[0])
            except (ValueError, IndexError):
                if row_number <= 20:
                    continue
                raise TableError('Esas No, Son Durum ve Muhatap Unvanı sütunları bulunamadı.')
            if len(row) < 3:
                raise TableError('Tabloda esas numarası, durum ve muhatap sütunları gerekli.')
            if len(row) > 3 and key(row[1]) not in _status_groups:
                raise TableError('Başlıksız tablonun sütunları doğrulanamadı. KKBS tablosunu sütun '
                                 'başlıklarıyla birlikte kopyalayın; ilk üç sütun esas numarası, '
                                 'durum ve muhatap olmalıdır.')
            mapping = {'number': 0, 'status': 1, 'recipient': 2}
        def cell(field):
            index = mapping.get(field)
            return row[index] if index is not None and index < len(row) else None
        try:
            number = parse_number(cell('number'))
        except ValueError:
            number = None
        records.append(Record(number, str(cell('status') or '').strip(),
                              str(cell('recipient') or '').strip(), term_name(cell('term')), row_number))
        if len(records) > MAX_ROWS:
            raise TableError('Tablo en fazla 100.000 veri satırı içerebilir.')
    if not records:
        raise TableError('Veri bulunamadı. Esas No, Son Durum ve Muhatap Unvanı sütunlarını kontrol edin.')
    if not any(record.number is not None for record in records):
        raise TableError('Esas No sütununda geçerli numara bulunamadı. 49220 veya 7/49220 biçiminde numaralar kullanılmalı.')
    return records


@contextmanager
def _excel_archive(data):
    if len(data) > MAX_TABLE_BYTES:
        raise TableError('Excel dosyası en fazla 12 MB olabilir.')
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(item.file_size for item in archive.infolist()) > MAX_UNPACKED_BYTES:
                raise TableError('Excel dosyasının açılmış boyutu çok büyük; ilgili satırları ayrı bir dosyaya alın.')
            yield archive
    except TableError:
        raise
    except Exception as exc:
        raise TableError('Excel okunamadı. Geçerli ve şifresiz bir .xlsx dosyası yükleyin.') from exc


def _workbook(data):
    with _excel_archive(data):
        try:
            return load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
        except Exception as exc:
            raise TableError('Excel okunamadı. Geçerli ve şifresiz bir .xlsx dosyası yükleyin.') from exc


def sheet_names(data):
    # Listing sheets must not load all shared strings in a large KKBS report.
    with _excel_archive(data) as archive:
        if 'xl/workbook.xml' in archive.namelist():
            root = ElementTree.fromstring(archive.read('xl/workbook.xml'))
            names = [sheet.attrib['name'] for sheet in root.findall('{*}sheets/{*}sheet')]
            if not names:
                raise TableError('Excel dosyasında çalışma sayfası bulunamadı.')
            return names
    # Unusual but valid workbooks may store their workbook part at another path.
    book = _workbook(data)
    try:
        return book.sheetnames
    finally:
        book.close()


def read_excel(data, sheet):
    book = _workbook(data)
    try:
        if sheet not in book.sheetnames:
            raise TableError('Seçilen çalışma sayfası bulunamadı.')
        worksheet = book[sheet]
        # Some exporters incorrectly declare a 1x1 used range; read the actual XML.
        worksheet.reset_dimensions()
        return read_records(worksheet.iter_rows(values_only=True))
    except TableError:
        raise
    except Exception as exc:
        raise TableError('Excel satırları okunamadı; dosyanın sağlam olduğunu kontrol edin.') from exc
    finally:
        book.close()


def read_paste(text):
    if len(text.encode('utf-8')) > MAX_TABLE_BYTES:
        raise TableError('Yapıştırılan tablo en fazla 12 MB olabilir.')
    try:
        delimiter = csv.Sniffer().sniff(text[:65536], delimiters='\t;,').delimiter
    except csv.Error:
        first_line = next((line for line in text.splitlines() if line.strip()), '')
        delimiter = '\t' if '\t' in first_line else ';' if ';' in first_line else ','
    try:
        return read_records(csv.reader(io.StringIO(text), delimiter=delimiter, strict=True))
    except csv.Error as exc:
        raise TableError('Tablo okunamadı; Excel’den satırları sütunlarıyla birlikte kopyalayın.') from exc


def number_ranges(numbers):
    numbers = sorted(set(numbers))
    ranges = []
    for number in numbers:
        if ranges and number == ranges[-1][1] + 1:
            ranges[-1][1] = number
        else:
            ranges.append([number, number])
    return format_ranges(ranges)


def format_ranges(ranges):
    # KKBS requires ASCII hyphens and both ends even for a single proposal.
    return [f'{start}-{end}' for start, end in ranges]


def add_batch(batches, report):
    """Save each query's ranges verbatim without mixing workflows or terms."""
    if not report['selected_count']:
        raise TableError('Bu sorguda birleşik listeye eklenecek önerge yok.')
    batch = {field: report[field] for field in
             ('stage', 'term', 'first', 'last', 'assessment_complete', 'report_version')}
    batch['source'] = report.get('source', 'Tablo')
    batch['ranges'] = list(report['ranges'])
    batch['numbers'] = [row['number'] for row in report['rows']
                        if row['group'] == WORKFLOW_STAGES[report['stage']]['group']]
    batch['id'] = f"{report['term']}:{report['first']}:{report['last']}"
    updated = []
    replaced = False
    for old in batches:
        if old['stage'] != batch['stage'] or old['term'] != batch['term']:
            raise TableError('Farklı işlem aşamalarına veya dönem bilgilerine ait sorgular birleştirilemez. '
                             'Önce mevcut birleşik listeyi indirin ve temizleyin.')
        if old['id'] == batch['id']:
            updated.append(batch)
            replaced = True
            continue
        first, last = max(old['first'], batch['first']), min(old['last'], batch['last'])
        if first <= last:
            old_numbers = {n for n in old['numbers'] if first <= n <= last}
            new_numbers = {n for n in batch['numbers'] if first <= n <= last}
            if old_numbers != new_numbers:
                raise TableError(f"{first}-{last} aralığı, eklenmiş bir sorguyla farklı sonuç veriyor. "
                                 'Eski sorguyu birleşik listeden çıkarıp güncel sonucu ekleyin.')
        updated.append(old)
    if not replaced:
        updated.append(batch)
    return updated


def combined_batch_ranges(batches):
    """Append saved query outputs in insertion order, preserving every boundary."""
    return [interval for batch in batches for interval in batch['ranges']]


def absent_ranges(present, first, last):
    intervals = []
    cursor = first
    for number in sorted(set(present)):
        if number > cursor:
            intervals.append([cursor, number-1])
        cursor = number+1
    if cursor <= last:
        intervals.append([cursor, last])
    return format_ranges(intervals)


def recipient_names(recipients):
    return [title[:-6] + 'Bakanlığı' if title.endswith('Bakanı') else title for title in recipients]


def infer_term(records, first, last):
    """Use source metadata without asking for a term or merging different terms."""
    terms = {record.term for record in records}
    if len(terms) == 1:
        return next(iter(terms))
    matching_terms = {record.term for record in records
                      if record.number is not None and first <= record.number <= last}
    if len(matching_terms) == 1:
        return next(iter(matching_terms))
    raise TableError('Bu aralık için kaynak dönem bilgisi tek bir döneme indirgenemedi. '
                     'İşlem yapmak istediğiniz dönemin kayıtlarını ayrı bir tablo olarak yükleyin.')


def build_report(records, first, last, term, recipients, *, stage='prepare'):
    if stage not in WORKFLOW_STAGES:
        raise TableError('Memur işlemi veya uzman kontrolü aşamasını seçin.')
    first, last = parse_number(first), parse_number(last)
    if first > last:
        raise TableError('Başlangıç esas numarası bitiş numarasından büyük olamaz.')
    if term is None:
        term = infer_term(records, first, last)
    terms = {record.term for record in records}
    if term not in terms:
        raise TableError('Kaynakta belirtilen döneme ait kayıt bulunamadı.')
    reference = recipient_names(recipients)
    aliases = {}
    for title, display in zip(recipients, reference):
        aliases[key(title)] = aliases[key(display)] = display
        if title.endswith('Bakanı'):
            aliases[key(title + 'na')] = aliases[key(display + 'na')] = display
        elif title == 'Cumhurbaşkanı Yardımcısı':
            aliases[key('Cumhurbaşkanı Yardımcılığı')] = display
    grouped = defaultdict(list)
    invalid_rows = []
    source_numbers = []
    for record in records:
        if record.term != term:
            continue
        if record.number is None:
            invalid_rows.append(record.source_row)
        else:
            source_numbers.append(record.number)
            if first <= record.number <= last:
                grouped[record.number].append(record)
    outcomes, duplicates = [], 0
    for number, items in sorted(grouped.items()):
        semantic = {(_status_groups.get(key(row.status), 'unknown'),
                     key(row.status) if key(row.status) not in _status_groups else '',
                     key(aliases.get(key(row.recipient), row.recipient))) for row in items}
        conflict = len(semantic) != 1
        duplicates += len(items)-1 if not conflict else 0
        row = items[0]
        outcomes.append({'number': number,
                         'group': 'conflict' if conflict else _status_groups.get(key(row.status), 'unknown'),
                         'status': ' / '.join(dict.fromkeys(x.status or '(boş)' for x in items)),
                         'recipient': ' / '.join(dict.fromkeys(aliases.get(key(x.recipient), x.recipient) or '(boş)' for x in items)),
                         'source_rows': [x.source_row for x in items]})
    selected = [row for row in outcomes if row['group'] == WORKFLOW_STAGES[stage]['group']]
    counts = Counter(row['group'] for row in outcomes)
    observed_recipients = Counter(row['recipient'] for row in selected)
    present_recipients = [name for name in reference if name in observed_recipients]
    missing_recipients = [name for name in reference if name not in observed_recipients]
    other_recipients = [name for name in observed_recipients if name not in reference]
    ranges = number_ranges(row['number'] for row in selected)
    absent_count = last-first+1-len(outcomes)
    workflow_warning = ''
    if stage == 'verify' and counts['pending']:
        workflow_warning = (f"{counts['pending']} önerge hâlâ “15 Günlük Cevaplanma Süresi Devam Ediyor” "
                            'durumunda. İlan girişlerinin tamamlanıp tamamlanmadığını kontrol edin; '
                            'bu kayıtlar ilan edilmiş listesine eklenmedi.')
    complete = not (absent_count or invalid_rows or counts['unknown'] or counts['conflict']
                    or other_recipients or workflow_warning)
    return {'report_version': REPORT_VERSION, 'stage': stage, 'first': first, 'last': last, 'term': term, 'counts': dict(counts),
            'source_first': min(source_numbers, default=None), 'source_last': max(source_numbers, default=None),
            'assessment_complete': complete,
            'workflow_warning': workflow_warning,
            'unrecognized_statuses': dict(Counter(row['status'] for row in outcomes if row['group'] == 'unknown')),
            'total': len(outcomes), 'selected_count': len(selected),
            'ranges': ranges, 'ranges_text': ';'.join(ranges),
            'ranges_by_status': {group: number_ranges(row['number'] for row in outcomes if row['group'] == group)
                                 for group in STATUS_LABELS},
            'present_recipients': present_recipients, 'missing_recipients': missing_recipients,
            'other_recipients': other_recipients, 'recipient_counts': dict(observed_recipients),
            'absent_count': absent_count,
            'absent_ranges': absent_ranges(grouped, first, last),
            'invalid_rows': invalid_rows, 'duplicate_count': duplicates, 'rows': outcomes}


def report_text(report):
    workflow = WORKFLOW_STAGES[report['stage']]
    lines = [f"Gelen Kağıt İşlemleri · {report['first']}-{report['last']}", workflow['label'],
             'Seçilen kaynak durumu: ' + STATUS_LABELS[workflow['group']], '',
             workflow['ranges_label'] + ':',
             report['ranges_text'] or 'Yok', f"Toplam: {report['selected_count']} önerge.", '']
    if not report['total']:
        lines.append('Seçilen aralıkta kaynak kaydı bulunamadı; bakanlıklar değerlendirilemedi.')
    elif not report['selected_count']:
        lines.append('Tanınan durumlara göre bu aralıkta seçilen aşamaya ait kayıt yok.')
    if not report['assessment_complete']:
        lines.append('VERİ KONTROLÜ GEREKLİ: veri veya işlem aşaması uyarıları giderilmeden sonuç tam liste olarak kabul edilmemeli.')
    if report['workflow_warning']:
        lines.append(report['workflow_warning'])
    for group, label in STATUS_LABELS.items():
        lines.append(f"{label}: {report['counts'].get(group, 0)}")
    if report['selected_count']:
        lines.extend(['', f"Bulunan bakanlıklar ve muhataplar ({len(report['present_recipients'])}):",
                      *report['present_recipients']])
        if report['assessment_complete'] or report['present_recipients']:
            lines.extend(['', f"Listede bulunmayanlar ({len(report['missing_recipients'])}):", *report['missing_recipients']])
            if not report['assessment_complete']:
                lines.append('Bu adlar yalnızca doğrulanabilen sonuçlarda görülmedi; kaynak eksikleri giderilmeden yoklukları kesinleştirilemez.')
    if report['unrecognized_statuses']:
        lines.extend(['', 'Tanınmayan durumlar:',
                      *(f'{status}: {count}' for status, count in report['unrecognized_statuses'].items())])
    if report['other_recipients']:
        lines.extend(['', 'Referans listesi dışında kalan muhataplar:', *report['other_recipients']])
    if report['absent_count']:
        lines.extend(['', f"Dosyada bulunmayan {report['absent_count']} esas numarası (sonuca dahil edilmedi):", *report['absent_ranges']])
    if report['invalid_rows']:
        lines.append('\nEsas numarası okunamayan kaynak satırları: ' + ', '.join(map(str, report['invalid_rows'])))
    return '\n'.join(lines)


def report_csv(report):
    output = io.StringIO()
    writer = csv.writer(output, delimiter=';')
    writer.writerow(['Esas No', 'Dönem', 'İşlem grubu', 'Kaynak durum', 'Muhatap', 'Kaynak satırları',
                     'Seçilen aşama', 'Aşama listesine dahil'])
    def safe_cell(value):
        text = str(value)
        return "'" + text if text.lstrip().startswith(('=', '+', '-', '@')) else text
    for row in report['rows']:
        writer.writerow([row['number'], safe_cell(report['term']), STATUS_LABELS[row['group']],
                         safe_cell(row['status']), safe_cell(row['recipient']), ', '.join(map(str, row['source_rows'])),
                         WORKFLOW_STAGES[report['stage']]['label'],
                         'Evet' if row['group'] == WORKFLOW_STAGES[report['stage']]['group'] else 'Hayır'])
    return output.getvalue().encode('utf-8-sig')
