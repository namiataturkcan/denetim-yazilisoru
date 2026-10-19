"""Build minimal reference data from the user's workbook; never import minister history.
Run with Python + openpyxl: python scripts/import_reference_workbook.py INPUT.xlsx
"""
import hashlib
import json
import statistics
import sys
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


def build(path, output):
    sheet = load_workbook(path, read_only=True, data_only=True).active
    rows = sheet.iter_rows(values_only=True)
    headers = next(rows)
    pairs = defaultdict(Counter)
    summaries = []
    terms = Counter()
    row_count = 0
    for values in rows:
        row = dict(zip(headers, values))
        if not row.get('Soru Sahibi'):
            continue
        row_count += 1
        terms[row['Donem Adı']] += 1
        if row.get('Soru Sahibi İli'):
            pairs[row['Soru Sahibi'].strip()][row['Soru Sahibi İli'].strip()] += 1
        if row.get('Özet'):
            summaries.append(' '.join(row['Özet'].split()))
    if set(terms) != {'28.DÖNEM'}:
        raise ValueError('Dönemler karışık; kaynak ayrı ayrı incelenmeli.')
    metadata = {'file': path.name, 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
                'term': '28.DÖNEM', 'row_count': row_count, 'member_count': len(pairs)}
    members = {name: next(iter(cities)) for name, cities in sorted(pairs.items()) if len(cities) == 1}
    conflicts = {name: dict(cities) for name, cities in pairs.items() if len(cities) != 1}
    reference = {'source': metadata, 'members': members, 'conflicts': conflicts,
                 'scope': 'Yüklenen önerge dökümünde yer alan 28. dönem eşleşmeleri; tüm milletvekillerinin veya aktif görev durumlarının listesi değildir.',
                 'invalid_constituency_labels': ['Dersim']}
    lengths = sorted(len(s.split()) for s in summaries)
    endings = ['iddiasına ilişkin', 'talebine ilişkin', 'çalışmalara ilişkin', 'sayısına ilişkin',
               'akıbetine ilişkin', 'sorununa ilişkin', 'önlemlere ilişkin', 'verilere ilişkin',
               'etkilerine ilişkin', 'durumuna ilişkin', 'ilişkin']
    examples = []
    # Diverse, deterministic examples; no raw workbook or full corpus is published.
    for ending in endings:
        pool = [s for s in summaries if s.endswith(ending) and 6 <= len(s.split()) <= 18 and s not in examples]
        for index in (len(pool)//3, 2*len(pool)//3):
            if pool and pool[index] not in examples:
                examples.append(pool[index])
    style = {'source': metadata, 'statistics': {'count': len(summaries), 'median_words': statistics.median(lengths),
        'p90_words': lengths[int(len(lengths)*.9)], 'max_words': max(lengths), 'over_10_words': sum(n > 10 for n in lengths)},
        'target_words': [6, 14], 'max_words': max(lengths), 'ending': 'ilişkin',
        'guidance': ['Tek, kısa bir konu tamlaması yaz; soruları tek tek sıralama.',
          'İddiaları olgu gibi sunma; metne uygunsa iddiasına/iddialara ilişkin kullan.',
          'Talep, sayı, akıbet, çalışma veya sorun ifadelerini belgedeki asıl konuya göre seç.',
          'Yer, kurum ve tarih yalnızca konuyu ayırt etmek için gerekliyse yer alsın.',
          'Örneklerden yeni bilgi veya konu kopyalama; yalnızca üslubu örnek al.',
          'Nötr bir dil kullan; değerlendirme, öneri veya cevap ekleme.'], 'examples': examples}
    for filename, value in [('member_reference.json', reference), ('summary_style.json', style)]:
        (output/filename).write_text(json.dumps(value, ensure_ascii=False, indent=2)+'\n')
    print(f'{row_count} records; {len(members)} unambiguous members; {len(conflicts)} conflicts; {len(examples)} style examples')


if __name__ == '__main__':
    build(Path(sys.argv[1]), Path(__file__).resolve().parents[1])
